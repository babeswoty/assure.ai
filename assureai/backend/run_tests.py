#!/usr/bin/env python
"""Standalone test runner for Excel import functionality."""

import sys
import traceback
from io import BytesIO
import openpyxl

# Add app to path
sys.path.insert(0, '.')

from app.excel_import import extract_racm_from_excel, ExcelValidationError, ExcelParsingError

def create_test_workbook(headers, rows, sheet_name="RACM Upload"):
    """Helper to create test Excel workbooks."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Write headers
    for header_name, col_idx in headers.items():
        ws.cell(row=1, column=col_idx, value=header_name)

    # Write data rows
    for row_idx, row_data in enumerate(rows, start=2):
        for header_name, col_idx in headers.items():
            value = row_data.get(header_name)
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def test_minimal_valid_racm():
    """Test importing minimal valid RACM with required fields."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Process Name": 3,
        "Risk Description": 4,
    }
    rows = [
        {
            "Control Ref": "A.C.C.1",
            "Risk Ref": "A.C.R.1",
            "Process Name": "User Access Management",
            "Risk Description": "Unauthorized access",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)
    result = extract_racm_from_excel(file_bytes)

    assert result["control"]["control_id"] == "A.C.C.1", f"Expected A.C.C.1, got {result['control']['control_id']}"
    assert result["control"]["risk"]["risk_id"] == "A.C.R.1", f"Expected A.C.R.1, got {result['control']['risk']['risk_id']}"
    assert result["engagement"]["process"] == "User Access Management"
    assert result["control"]["risk"]["risk_description"] == "Unauthorized access"
    print("✓ test_minimal_valid_racm PASSED")

def test_all_fields_populated():
    """Test importing all available RACM fields."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Mega Process": 3,
        "Major Process Name": 4,
        "Process Name": 5,
        "Sub-process": 6,
        "Control Objective": 7,
        "Risk Description": 8,
        "AML/CTF Control": 9,
        "Suggested Control": 10,
        "Control Nature": 11,
        "Control Type": 12,
        "Frequency of Control": 13,
        "Control Evidence": 14,
    }
    rows = [
        {
            "Control Ref": "CTRL-001",
            "Risk Ref": "RISK-001",
            "Mega Process": "IT Controls",
            "Major Process Name": "Access Management",
            "Process Name": "User Provisioning",
            "Sub-process": "Onboarding",
            "Control Objective": "Ensure authorized access",
            "Risk Description": "Unauthorized system access",
            "AML/CTF Control": "Access Control",
            "Suggested Control": "Review and approve all access requests",
            "Control Nature": "Detective",
            "Control Type": "Automated",
            "Frequency of Control": "Monthly",
            "Control Evidence": "Access logs and approval records",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)
    result = extract_racm_from_excel(file_bytes)

    assert result["control"]["control_id"] == "CTRL-001"
    assert result["control"]["control_title"] == "Access Control"
    assert result["engagement"]["mega_process"] == "IT Controls"
    assert result["engagement"]["area"] == "Access Management"
    assert result["control"]["suggested_control_description"] == "Review and approve all access requests"
    print("✓ test_all_fields_populated PASSED")

def test_multiline_cells_preserved():
    """Test that multiline cells preserve line breaks."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Process Name": 3,
        "Risk Description": 4,
        "Suggested Control": 5,
        "Control Evidence": 6,
    }
    rows = [
        {
            "Control Ref": "A.C.C.1",
            "Risk Ref": "A.C.R.1",
            "Process Name": "Process",
            "Risk Description": "Risk",
            "Suggested Control": "Step 1\nStep 2\nStep 3",
            "Control Evidence": "Point 1\nPoint 2\nPoint 3",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)
    result = extract_racm_from_excel(file_bytes)

    assert "\n" in result["control"]["suggested_control_description"], "Multiline not preserved in Suggested Control"
    assert "\n" in result["control"]["expected_evidence"], "Multiline not preserved in Control Evidence"
    print("✓ test_multiline_cells_preserved PASSED")

def test_header_aliases():
    """Test that header aliases work."""
    headers = {
        "control ref (#)": 1,
        "risk ref (#)": 2,
        "process name": 3,
        "risk description": 4,
    }
    rows = [
        {
            "control ref (#)": "A.C.C.1",
            "risk ref (#)": "A.C.R.1",
            "process name": "Process",
            "risk description": "Risk",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)
    result = extract_racm_from_excel(file_bytes)

    assert result["control"]["control_id"] == "A.C.C.1"
    assert result["control"]["risk"]["risk_id"] == "A.C.R.1"
    print("✓ test_header_aliases PASSED")

def test_blank_optional_fields():
    """Test that blank optional fields are handled correctly."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Process Name": 3,
        "Risk Description": 4,
        "Control Objective": 5,
        "Control Nature": 6,
    }
    rows = [
        {
            "Control Ref": "A.C.C.1",
            "Risk Ref": "A.C.R.1",
            "Process Name": "Process",
            "Risk Description": "Risk",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)
    result = extract_racm_from_excel(file_bytes)

    assert result["control"]["control_objective"] is None
    assert result["control"]["control_nature"] is None
    print("✓ test_blank_optional_fields PASSED")

def test_missing_control_ref_fails():
    """Test that missing Control Ref returns validation error."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Process Name": 3,
        "Risk Description": 4,
    }
    rows = [
        {
            "Risk Ref": "A.C.R.1",
            "Process Name": "Process",
            "Risk Description": "Risk",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)

    try:
        extract_racm_from_excel(file_bytes)
        print("✗ test_missing_control_ref_fails FAILED - Expected ExcelValidationError")
        return False
    except ExcelValidationError as e:
        assert len(e.errors) > 0
        assert any(err["field"] == "control.control_id" for err in e.errors)
        print("✓ test_missing_control_ref_fails PASSED")
        return True

def test_missing_risk_ref_fails():
    """Test that missing Risk Ref returns validation error."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Process Name": 3,
        "Risk Description": 4,
    }
    rows = [
        {
            "Control Ref": "A.C.C.1",
            "Process Name": "Process",
            "Risk Description": "Risk",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)

    try:
        extract_racm_from_excel(file_bytes)
        print("✗ test_missing_risk_ref_fails FAILED - Expected ExcelValidationError")
        return False
    except ExcelValidationError as e:
        assert len(e.errors) > 0
        assert any(err["field"] == "control.risk.risk_id" for err in e.errors)
        print("✓ test_missing_risk_ref_fails PASSED")
        return True

def test_default_values_applied():
    """Test that safe defaults are applied for missing engagement fields."""
    headers = {
        "Control Ref": 1,
        "Risk Ref": 2,
        "Process Name": 3,
        "Risk Description": 4,
    }
    rows = [
        {
            "Control Ref": "A.C.C.1",
            "Risk Ref": "A.C.R.1",
            "Process Name": "Process",
            "Risk Description": "Risk",
        }
    ]

    file_bytes = create_test_workbook(headers, rows)
    result = extract_racm_from_excel(file_bytes)

    assert result["engagement"]["client"] == "MTN Fintech"
    assert result["engagement"]["opco"] == "Group Fintech / Fintech OpCo"
    assert result["engagement"]["process"] == "Process"
    print("✓ test_default_values_applied PASSED")

def test_racm_upload_worksheet_preferred():
    """Test that RACM Upload worksheet is used if it exists."""
    wb = openpyxl.Workbook()
    wb.create_sheet("Read Me")
    racm_sheet = wb.create_sheet("RACM Upload")

    racm_sheet.cell(row=1, column=1, value="Control Ref")
    racm_sheet.cell(row=1, column=2, value="Risk Ref")
    racm_sheet.cell(row=1, column=3, value="Process Name")
    racm_sheet.cell(row=1, column=4, value="Risk Description")

    racm_sheet.cell(row=2, column=1, value="A.C.C.1")
    racm_sheet.cell(row=2, column=2, value="A.C.R.1")
    racm_sheet.cell(row=2, column=3, value="Process")
    racm_sheet.cell(row=2, column=4, value="Risk")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    file_bytes = output.getvalue()

    result = extract_racm_from_excel(file_bytes)
    assert result["control"]["control_id"] == "A.C.C.1"
    print("✓ test_racm_upload_worksheet_preferred PASSED")

def test_header_detection_across_rows():
    """Test that header row is detected even if not in row 1."""
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.cell(row=1, column=1, value="Title Row")
    ws.cell(row=2, column=1, value="Subtitle Row")

    ws.cell(row=3, column=1, value="Control Ref")
    ws.cell(row=3, column=2, value="Risk Ref")
    ws.cell(row=3, column=3, value="Process Name")
    ws.cell(row=3, column=4, value="Risk Description")

    ws.cell(row=4, column=1, value="A.C.C.1")
    ws.cell(row=4, column=2, value="A.C.R.1")
    ws.cell(row=4, column=3, value="Process")
    ws.cell(row=4, column=4, value="Risk")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    file_bytes = output.getvalue()

    result = extract_racm_from_excel(file_bytes)
    assert result["control"]["control_id"] == "A.C.C.1"
    print("✓ test_header_detection_across_rows PASSED")

def test_empty_workbook_fails():
    """Test that empty workbook returns parsing error."""
    wb = openpyxl.Workbook()
    ws = wb.active

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    file_bytes = output.getvalue()

    try:
        extract_racm_from_excel(file_bytes)
        print("✗ test_empty_workbook_fails FAILED - Expected ExcelParsingError")
        return False
    except ExcelParsingError:
        print("✓ test_empty_workbook_fails PASSED")
        return True

if __name__ == "__main__":
    print("=" * 70)
    print("RUNNING EXCEL IMPORT TEST SUITE")
    print("=" * 70)

    tests = [
        test_minimal_valid_racm,
        test_all_fields_populated,
        test_multiline_cells_preserved,
        test_header_aliases,
        test_blank_optional_fields,
        test_missing_control_ref_fails,
        test_missing_risk_ref_fails,
        test_default_values_applied,
        test_racm_upload_worksheet_preferred,
        test_header_detection_across_rows,
        test_empty_workbook_fails,
    ]

    passed = 0
    failed = 0

    for test in tests:
        try:
            result = test()
            if result is None or result is True:
                passed += 1
            else:
                failed += 1
        except Exception as e:
            failed += 1
            print(f"✗ {test.__name__} FAILED with exception:")
            traceback.print_exc()

    print()
    print("=" * 70)
    print(f"RESULTS: {passed} passed, {failed} failed out of {len(tests)} tests")
    print("=" * 70)

    sys.exit(0 if failed == 0 else 1)
