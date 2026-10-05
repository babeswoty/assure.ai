from io import BytesIO
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from app.excel_import import STANDARD_HEADERS
from app.main import app

client = TestClient(app)
SAMPLES = Path(__file__).resolve().parents[2] / "samples"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _upload_racm() -> dict:
    data = (SAMPLES / "sample_racm.xlsx").read_bytes()
    r = client.post("/api/convert-excel-to-racm", files={"file": ("sample_racm.xlsx", data, XLSX_MIME)})
    assert r.status_code == 200, r.text
    return r.json()


def test_sample_racm_uses_exactly_the_standard_headers():
    ws = openpyxl.load_workbook(SAMPLES / "sample_racm.xlsx")["RACM Upload"]
    assert [c.value for c in ws[1]] == STANDARD_HEADERS


def test_samples_contain_no_real_company_names():
    text = (SAMPLES / "sample_walkthrough_transcript.txt").read_text(encoding="utf-8")
    workbook = openpyxl.load_workbook(SAMPLES / "sample_racm.xlsx")
    for sheet in workbook:
        text += " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    for name in ("MTN", "Deloitte", "Fintech"):
        assert name.lower() not in text.lower()


def test_sample_racm_uploads_into_a_valid_record():
    record = _upload_racm()
    assert record["control"]["control_id"] == "SAMPLE-AP-001"
    assert record["control"]["control_reference"] == "FIN-AP-01"
    assert record["engagement"]["client"].startswith("Sample Client")
    assert client.post("/api/records/validate", json={"record": record}).status_code == 200


def test_sample_transcript_txt_and_docx_give_the_same_text():
    txt = (SAMPLES / "sample_walkthrough_transcript.txt").read_bytes()
    docx = (SAMPLES / "sample_walkthrough_transcript.docx").read_bytes()
    from_txt = client.post("/api/extract-text-from-document", files={"file": ("t.txt", txt, "text/plain")}).json()["text"]
    from_docx = client.post("/api/extract-text-from-document", files={"file": ("t.docx", docx, DOCX_MIME)}).json()["text"]
    assert from_txt.strip() == from_docx.strip()


def test_samples_run_through_the_whole_pipeline_to_an_adequate_control_and_an_export(monkeypatch):
    monkeypatch.delenv("ASSUREAI_CONDENSE")  # run the way the app does: condensing on
    record = _upload_racm()
    questions = client.post("/api/questions/generate", json={"record": record}).json()
    transcript = (SAMPLES / "sample_walkthrough_transcript.txt").read_text(encoding="utf-8")

    extraction = client.post(
        "/api/transcripts/extract",
        json={"record": record, "questions": questions, "transcript_text": transcript},
    ).json()
    doc = extraction["documentation"]
    for attribute in ("what", "who", "when", "why", "where", "how"):
        assert doc[attribute]["value"], f"{attribute} was not documented"
        assert doc[attribute]["source_references"], f"{attribute} has no source"
    assert doc["missing_information"] == []

    record["control"]["documentation"] = doc
    design = client.post("/api/design-assessment/assess", json={"control": record["control"]}).json()
    assert design["rating_label"] == "Adequate"

    gaps = client.post(
        "/api/gaps/draft", json={"control": record["control"], "design_assessment": design["design_assessment"]}
    ).json()
    assert gaps["gaps"] == []

    record["control"]["design_assessment"] = design["design_assessment"]
    record["control"]["gap_assessment"] = gaps
    record["control"]["review"] = {
        "decision": "approved", "reviewer_id": "j.smith", "reviewed_at": "2026-01-01T00:00:00Z",
        "human_approval_required": True,
    }
    exported = client.post("/api/export/excel", json={"record": record})
    assert exported.status_code == 200
    sheet = openpyxl.load_workbook(BytesIO(exported.content)).active
    headers = {sheet.cell(row=2, column=i).value: i for i in range(1, sheet.max_column + 1)}
    assert sheet.cell(row=3, column=headers["Control ID"]).value == "SAMPLE-AP-001"
    assert sheet.cell(row=3, column=headers["Applicable"]).value == "Yes"


# ---------- the sample with a design problem: Inadequate, with a gap ----------


def _run_inadequate_sample():
    data = (SAMPLES / "sample_inadequate_racm.xlsx").read_bytes()
    record = client.post("/api/convert-excel-to-racm", files={"file": ("r.xlsx", data, XLSX_MIME)}).json()
    questions = client.post("/api/questions/generate", json={"record": record}).json()
    transcript = (SAMPLES / "sample_inadequate_walkthrough_transcript.txt").read_text(encoding="utf-8")
    doc = client.post(
        "/api/transcripts/extract",
        json={"record": record, "questions": questions, "transcript_text": transcript},
    ).json()["documentation"]
    record["control"]["documentation"] = doc
    design = client.post("/api/design-assessment/assess", json={"control": record["control"]}).json()
    gaps = client.post(
        "/api/gaps/draft", json={"control": record["control"], "design_assessment": design["design_assessment"]}
    ).json()
    return record, doc, design, gaps


def test_inadequate_sample_documents_all_six_attributes_in_short_phrases(monkeypatch):
    monkeypatch.delenv("ASSUREAI_CONDENSE")  # as the app runs
    _, doc, _, _ = _run_inadequate_sample()
    for attribute in ("what", "who", "when", "why", "where", "how"):
        assert doc[attribute]["value"], f"{attribute} was not documented"
    assert doc["when"]["value"] == "Monthly; triggered whenever a correction is needed"
    assert doc["why"]["value"] == "To prevent errors and unauthorised entries in the general ledger"
    assert doc["where"]["value"] == "Finance System"
    assert doc["missing_information"] == []


def test_inadequate_sample_is_rated_inadequate_because_the_same_person_prepares_and_approves(monkeypatch):
    monkeypatch.delenv("ASSUREAI_CONDENSE")
    _, _, design, _ = _run_inadequate_sample()
    assert design["rating_label"] == "Inadequate"
    assert design["design_assessment"]["completeness"]["result"] == "pass"
    assert "independence" in design["design_assessment"]["conclusion_rationale"].lower()
    assert "i'm also the one who approves" in design["design_assessment"]["conclusion_rationale"].lower()


def test_inadequate_sample_drafts_a_gap_and_exports_it_with_applicable_no(monkeypatch):
    monkeypatch.delenv("ASSUREAI_CONDENSE")
    record, _, design, gaps = _run_inadequate_sample()
    assert len(gaps["gaps"]) == 1
    gap = gaps["gaps"][0]
    assert gap["gap_title"] == "Lack of independence between control roles"
    assert gap["root_cause_category"] == "Segregation of duties"
    assert gap["recommendation"]

    record["control"]["design_assessment"] = design["design_assessment"]
    record["control"]["gap_assessment"] = gaps
    record["control"]["review"] = {
        "decision": "approved", "reviewer_id": "j.smith", "reviewed_at": "2026-01-01T00:00:00Z",
        "human_approval_required": True,
    }
    exported = client.post("/api/export/excel", json={"record": record})
    assert exported.status_code == 200
    sheet = openpyxl.load_workbook(BytesIO(exported.content)).active
    cols = {sheet.cell(row=2, column=i).value: i for i in range(1, sheet.max_column + 1)}
    cell = lambda header: sheet.cell(row=3, column=cols[header]).value

    assert cell("Control ID") == "SAMPLE-GL-001"
    assert cell("Applicable") == "No"
    assert "independence" in cell("Reason if Not Applicable").lower()
    assert cell("Gap title") == "Lack of independence between control roles"
    assert cell("RCA") == "Segregation of duties"
    assert cell("Gap") != "N/A" and cell("Action Plan") != "N/A"


def test_inadequate_sample_transcript_txt_and_docx_give_the_same_text():
    txt = (SAMPLES / "sample_inadequate_walkthrough_transcript.txt").read_bytes()
    docx = (SAMPLES / "sample_inadequate_walkthrough_transcript.docx").read_bytes()
    from_txt = client.post("/api/extract-text-from-document", files={"file": ("t.txt", txt, "text/plain")}).json()["text"]
    from_docx = client.post("/api/extract-text-from-document", files={"file": ("t.docx", docx, DOCX_MIME)}).json()["text"]
    assert from_txt.strip() == from_docx.strip()


def test_both_sample_racms_use_exactly_the_standard_headers():
    for name in ("sample_racm.xlsx", "sample_inadequate_racm.xlsx"):
        ws = openpyxl.load_workbook(SAMPLES / name)["RACM Upload"]
        assert [c.value for c in ws[1]] == STANDARD_HEADERS, name
