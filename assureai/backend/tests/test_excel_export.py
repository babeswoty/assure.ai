from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.racm_export import build_racm_excel_bytes, build_racm_workbook, is_approved_for_export
from app.racm_export.exporter import DOCUMENTATION_FIELDS, GAP_FIELDS, RACM_INPUT_FIELDS
from app.schemas.racm import (
    AssessmentDimension,
    AssessmentResult,
    Control,
    ControlReview,
    DesignAssessment,
    Engagement,
    Gap,
    GapAssessment,
    GapImpact,
    GapType,
    RACMControlRecord,
    RemediationAlignment,
    Risk,
    ReviewDecision,
)


def _record(
    decision: ReviewDecision = ReviewDecision.NOT_REVIEWED,
    design_conclusion: AssessmentResult = AssessmentResult.NOT_ASSESSED,
    conclusion_rationale: str | None = None,
    relevance_rationale: str | None = None,
    gaps: list | None = None,
) -> RACMControlRecord:
    control = Control(
        control_id="PILOT-A-UAM-001",
        control_title="New User Access Provisioning Approval",
        control_reference="ITGC-UAM-01",
        risk=Risk(risk_id="R-1", risk_description="Inappropriate access could be provisioned."),
        control_objective="To ensure that access requests are approved before provisioning.",
        suggested_control_description="Manager approves; IT Security provisions per role profile.",
        design_assessment=DesignAssessment(
            completeness=AssessmentDimension(result=AssessmentResult.PASS),
            relevance=AssessmentDimension(result=AssessmentResult.PASS, rationale=relevance_rationale),
            risk_objective_control_alignment=AssessmentDimension(result=AssessmentResult.PASS),
            first_pass_design_conclusion=design_conclusion,
            conclusion_rationale=conclusion_rationale,
        ),
        gap_assessment=GapAssessment(gaps=gaps or []),
        review=ControlReview(
            decision=decision,
            reviewer_id="jsmith" if decision != ReviewDecision.NOT_REVIEWED else None,
            reviewed_at=datetime.now(timezone.utc) if decision != ReviewDecision.NOT_REVIEWED else None,
        ),
    )
    control.documentation.what.value = "Manager submits a request."
    control.documentation.who.value = "IT Security reviews it."
    control.documentation.control_description = "What: Manager submits a request.\n\nWho: IT Security reviews it."
    return RACMControlRecord(
        engagement=Engagement(
            engagement_id="ENG-1",
            client="Solstice (SYNTHETIC)",
            process="User Access Management",
        ),
        control=control,
    )


def _gap(**overrides) -> Gap:
    defaults = dict(
        gap_id="PILOT-A-UAM-001-GAP-01",
        gap_type=GapType.DESIGN_GAP,
        gap_title="Absence of independent review",
        gap_description="The same individual both provisions and reviews privileged access.",
        root_cause_category="Segregation of duties",
        impact=GapImpact.HIGH,
        recommendation="Assign review to someone outside the provisioning process.",
        action_addresses_gap=RemediationAlignment.NOT_ASSESSED,
    )
    defaults.update(overrides)
    return Gap(**defaults)


def test_is_approved_for_export_true_for_approved_decisions():
    assert is_approved_for_export(_record(ReviewDecision.APPROVED)) is True
    assert is_approved_for_export(_record(ReviewDecision.APPROVED_WITH_EDITS)) is True


def test_is_approved_for_export_false_for_unreviewed_or_rejected():
    assert is_approved_for_export(_record(ReviewDecision.NOT_REVIEWED)) is False


def test_workbook_header_row_has_the_expected_field_names_in_order():
    workbook = build_racm_workbook(_record(ReviewDecision.APPROVED))
    sheet = workbook.active

    expected_headers = (
        [label for label, _, _ in RACM_INPUT_FIELDS]
        + [label for label, _, _ in DOCUMENTATION_FIELDS]
        + [label for label, _, _ in GAP_FIELDS]
    )
    actual_headers = [cell.value for cell in sheet[2]]
    assert actual_headers == expected_headers


def test_workbook_group_row_labels_and_merges_match_the_requested_layout():
    workbook = build_racm_workbook(_record(ReviewDecision.APPROVED))
    sheet = workbook.active
    header_row = {cell.value: cell.column for cell in sheet[2]}
    merged_ranges = list(sheet.merged_cells.ranges)

    def group_label_above(header: str) -> str:
        # Merging leaves only the top-left cell holding the value - read
        # that cell for any column the merge covers.
        col = header_row[header]
        for merged in merged_ranges:
            if merged.min_row == 1 and merged.min_col <= col <= merged.max_col:
                return sheet.cell(row=1, column=merged.min_col).value
        return sheet.cell(row=1, column=col).value

    assert group_label_above("Engagement ID") == "Engagement Details"
    assert group_label_above("Control ID") == "Control Details"
    assert group_label_above("Risk ID") == "Risk Details"
    assert group_label_above("Risk Description") == "Risk Details"
    assert group_label_above("Control Objective") == "Control Details"
    assert group_label_above("Applicable") == "Structured Control Documentation"
    assert group_label_above("What") == "Structured Control Documentation"
    assert group_label_above("Gap") == "Gap"
    assert group_label_above("Action Date") == "Gap"

    risk_id_col, risk_desc_col = header_row["Risk ID"], header_row["Risk Description"]
    assert any(
        m.min_row == 1 and m.max_row == 1 and m.min_col == risk_id_col and m.max_col == risk_desc_col
        for m in merged_ranges
    )


def test_workbook_fills_distinguish_the_three_column_groups():
    workbook = build_racm_workbook(_record(ReviewDecision.APPROVED))
    sheet = workbook.active
    racm_input_count = len(RACM_INPUT_FIELDS)
    documentation_end = racm_input_count + len(DOCUMENTATION_FIELDS)
    for cell in sheet[2][:racm_input_count]:
        assert cell.fill.start_color.rgb == "00FFA500"
    for cell in sheet[2][racm_input_count:documentation_end]:
        assert cell.fill.start_color.rgb == "0090EE90"
    for cell in sheet[2][documentation_end:]:
        assert cell.fill.start_color.rgb == "00C6E0B4"


def test_workbook_data_row_reflects_record_values():
    workbook = build_racm_workbook(_record(ReviewDecision.APPROVED))
    sheet = workbook.active
    header_row = {sheet.cell(row=2, column=i).value: i for i in range(1, sheet.max_column + 1)}
    row = {label: sheet.cell(row=3, column=col).value for label, col in header_row.items()}
    assert row["Control ID"] == "PILOT-A-UAM-001"
    assert row["Client"] == "Solstice (SYNTHETIC)"
    assert row["Control Objective"] == "To ensure that access requests are approved before provisioning."
    assert row["What"] == "Manager submits a request."
    assert "Draft Control Description" not in row


def _applicable_and_reason(record: RACMControlRecord) -> tuple:
    workbook = build_racm_workbook(record)
    sheet = workbook.active
    header_row = {sheet.cell(row=2, column=i).value: i for i in range(1, sheet.max_column + 1)}
    return (
        sheet.cell(row=3, column=header_row["Applicable"]).value,
        sheet.cell(row=3, column=header_row["Reason if Not Applicable"]).value,
    )


def test_adequate_design_rating_shows_applicable_yes():
    record = _record(ReviewDecision.APPROVED, design_conclusion=AssessmentResult.PASS)
    assert _applicable_and_reason(record) == ("Yes", "Yes")


def test_inadequate_design_rating_shows_applicable_no_with_the_design_assessment_reason():
    record = _record(
        ReviewDecision.APPROVED,
        design_conclusion=AssessmentResult.FAIL,
        conclusion_rationale="The control objective requires independence, but the same person performs both roles.",
    )
    assert _applicable_and_reason(record) == (
        "No",
        "The control objective requires independence, but the same person performs both roles.",
    )


def test_inadequate_reason_appends_the_relevance_rationale_when_it_has_one():
    record = _record(
        ReviewDecision.APPROVED,
        design_conclusion=AssessmentResult.FAIL,
        conclusion_rationale="The control objective requires independence, but the same person performs both roles.",
        relevance_rationale="Term overlap between the documented activity and the stated risk/objective is low (ratio 0.08).",
    )
    _, reason = _applicable_and_reason(record)
    assert reason == (
        "The control objective requires independence, but the same person performs both roles.\n\n"
        "Relevance: Term overlap between the documented activity and the stated risk/objective is low (ratio 0.08)."
    )


def test_inadequate_reason_falls_back_to_na_when_no_conclusion_rationale_is_available():
    record = _record(ReviewDecision.APPROVED, design_conclusion=AssessmentResult.FAIL, conclusion_rationale=None)
    assert _applicable_and_reason(record) == ("No", "N/A")


def test_unassessed_or_requires_human_review_leaves_both_columns_blank():
    not_assessed = _record(ReviewDecision.APPROVED, design_conclusion=AssessmentResult.NOT_ASSESSED)
    assert _applicable_and_reason(not_assessed) == (None, None)

    requires_review = _record(ReviewDecision.APPROVED, design_conclusion=AssessmentResult.UNKNOWN)
    assert _applicable_and_reason(requires_review) == (None, None)


def test_build_racm_excel_bytes_round_trips_through_openpyxl():
    content = build_racm_excel_bytes(_record(ReviewDecision.APPROVED))
    workbook = load_workbook(BytesIO(content))
    assert workbook.active.cell(row=1, column=1).value == "Engagement Details"
    assert workbook.active.cell(row=2, column=1).value == "Engagement ID"


def _gap_row(record: RACMControlRecord) -> dict:
    workbook = build_racm_workbook(record)
    sheet = workbook.active
    header_row = {sheet.cell(row=2, column=i).value: i for i in range(1, sheet.max_column + 1)}
    gap_labels = [label for label, _, _ in GAP_FIELDS]
    return {label: sheet.cell(row=3, column=header_row[label]).value for label in gap_labels}


def test_no_gap_drafted_shows_na_across_the_gap_columns():
    record = _record(ReviewDecision.APPROVED, gaps=[])
    row = _gap_row(record)
    assert row["Gap"] == "N/A"
    assert row["Gap title"] == "N/A"
    assert row["Reason"] == "N/A"
    assert row["RCA"] == "N/A"
    assert row["Action Plan"] == "N/A"
    assert row["Action Owner"] is None
    assert row["Action Date"] is None


def test_drafted_gap_populates_the_gap_columns_from_the_first_gap():
    gap = _gap(
        gap_title="Absence of independent review",
        gap_description="The same individual both provisions and reviews privileged access.",
        root_cause_category="Segregation of duties",
        recommendation="Assign review to someone outside the provisioning process.",
    )
    record = _record(ReviewDecision.APPROVED, gaps=[gap])
    row = _gap_row(record)
    assert row["Gap"] == "The same individual both provisions and reviews privileged access."
    assert row["Gap title"] == "Absence of independent review"
    assert row["Reason"] == (
        "To ensure that access requests are approved before provisioning.\n\n"
        "The same individual both provisions and reviews privileged access."
    )
    assert row["RCA"] == "Segregation of duties"
    assert row["Action Plan"] == "Assign review to someone outside the provisioning process."
    assert row["Action Owner"] is None
    assert row["Action Date"] is None


def test_multiple_gaps_stay_on_one_row_as_a_numbered_list_per_column():
    gap_one = _gap(
        gap_title="Absence of independent review",
        gap_description="The same individual both provisions and reviews privileged access.",
        root_cause_category="Segregation of duties",
        recommendation="Assign review to someone outside the provisioning process.",
    )
    gap_two = _gap(
        gap_id="PILOT-A-UAM-001-GAP-02",
        gap_title="No periodic recertification",
        gap_description="Access is never recertified after initial provisioning.",
        root_cause_category="Process deficiency",
        recommendation="Introduce a quarterly access recertification process.",
    )
    record = _record(ReviewDecision.APPROVED, gaps=[gap_one, gap_two])
    row = _gap_row(record)

    assert row["Gap"] == (
        "1. The same individual both provisions and reviews privileged access.\n\n"
        "2. Access is never recertified after initial provisioning."
    )
    assert row["Gap title"] == "1. Absence of independent review\n\n2. No periodic recertification"
    assert row["RCA"] == "1. Segregation of duties\n\n2. Process deficiency"
    assert row["Action Plan"] == (
        "1. Assign review to someone outside the provisioning process.\n\n"
        "2. Introduce a quarterly access recertification process."
    )
    assert row["Reason"] == (
        "To ensure that access requests are approved before provisioning.\n\n"
        "1. The same individual both provisions and reviews privileged access.\n\n"
        "2. Access is never recertified after initial provisioning."
    )
    # Still exactly one data row for the control - gaps never split into extra rows
    # (rows past 3 exist only because Action Date's number format/validation is
    # pre-applied for future manual rows; none of them hold a value).
    workbook = build_racm_workbook(record)
    sheet = workbook.active
    assert all(cell.value is None for row in sheet.iter_rows(min_row=4) for cell in row)


def test_action_date_column_gets_a_date_number_format_and_a_date_validation():
    workbook = build_racm_workbook(_record(ReviewDecision.APPROVED, gaps=[_gap()]))
    sheet = workbook.active
    header_row = {sheet.cell(row=2, column=i).value: i for i in range(1, sheet.max_column + 1)}
    action_date_col = header_row["Action Date"]
    col_letter = sheet.cell(row=3, column=action_date_col).column_letter

    assert sheet.cell(row=3, column=action_date_col).number_format == "yyyy-mm-dd"
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    assert validations[0].type == "date"
    assert f"{col_letter}3" in validations[0].sqref


def test_reviewer_filled_action_owner_and_date_pass_through_for_a_single_gap():
    gap = _gap(action_owner="J. Smith", action_date=date(2026, 3, 15))
    record = _record(ReviewDecision.APPROVED, gaps=[gap])
    row = _gap_row(record)
    assert row["Action Owner"] == "J. Smith"
    assert row["Action Date"] == date(2026, 3, 15)


def test_unfilled_action_owner_and_date_stay_blank_not_na():
    record = _record(ReviewDecision.APPROVED, gaps=[_gap()])
    row = _gap_row(record)
    assert row["Action Owner"] is None
    assert row["Action Date"] is None


def test_multiple_gaps_number_action_owner_and_date_as_text():
    gap_one = _gap(action_owner="J. Smith", action_date=date(2026, 3, 15))
    gap_two = _gap(gap_id="PILOT-A-UAM-001-GAP-02", action_owner=None, action_date=None)
    record = _record(ReviewDecision.APPROVED, gaps=[gap_one, gap_two])
    row = _gap_row(record)
    assert row["Action Owner"] == "1. J. Smith\n\n2. "
    assert row["Action Date"] == "1. 2026-03-15\n\n2. "
