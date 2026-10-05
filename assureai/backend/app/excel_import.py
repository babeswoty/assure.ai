"""Excel RACM importer: reads one control from a RACM spreadsheet.

The standard headers are the ones this app's own Excel export uses (see
STANDARD_HEADERS), so a RACM downloaded from the app can be uploaded again.
Common alternative names (for example "Risk Ref", "Process Name",
"Suggested Control") are accepted as aliases. Nothing client-specific is
assumed: a client, OpCo or engagement id is only used if the file has a
column for it.
"""

import html
import re
from io import BytesIO
from typing import Any, Dict, List, Optional

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl is a declared requirement
    openpyxl = None


class ExcelParsingError(Exception):
    """Base exception for Excel parsing issues."""


class ExcelValidationError(ExcelParsingError):
    """Exception for data validation issues with structured error details."""

    def __init__(self, errors: List[Dict[str, Any]]):
        self.errors = errors
        super().__init__(f"Excel validation failed with {len(errors)} error(s)")


# The app's standard RACM input headers, in column order.
STANDARD_HEADERS = [
    "Engagement ID",
    "Client",
    "OpCo",
    "Area",
    "Mega Process",
    "Process",
    "Sub-process",
    "Control ID",
    "Control Reference",
    "Control Title",
    "Risk ID",
    "Risk Description",
    "Control Objective",
    "Suggested Control Description",
    "Control Frequency",
    "Control Nature",
    "Control Type",
    "Expected Evidence",
    "Suggested Test Procedures (ToD)",
    "Suggested Test Procedures (ToE)",
]

# Other names people commonly use for the same columns.
_EXTRA_ALIASES = {
    "OpCo": ["operating company"],
    "Area": ["major process", "major process name"],
    "Process": ["process name"],
    "Sub-process": ["subprocess", "sub process"],
    "Control Reference": ["control ref", "control ref (#)"],
    "Control Title": ["control name"],
    "Risk ID": ["risk ref", "risk ref (#)", "risk reference", "risk refence"],
    "Suggested Control Description": ["suggested control"],
    "Control Frequency": ["frequency of control", "frequency"],
    "Control Type": ["control type (manual/automated/semi automated)"],
    "Expected Evidence": ["control evidence"],
    "Suggested Test Procedures (ToD)": ["suggested test procedures tod", "test procedures tod", "test procedures (tod)"],
    "Suggested Test Procedures (ToE)": ["suggested test procedures toe", "test procedures toe", "test procedures (toe)"],
}

# Fields whose line breaks are worth keeping (steps, bullet points).
_MULTILINE = {
    "Control Objective",
    "Suggested Control Description",
    "Expected Evidence",
    "Suggested Test Procedures (ToD)",
    "Suggested Test Procedures (ToE)",
}

_NO_VALUE = {"nan", "none", "undefined", ""}


def _normalize_header(text: Any) -> str:
    """Trim, decode HTML entities, turn line breaks into spaces, collapse
    repeated spaces, lowercase."""
    if text is None:
        return ""
    text = html.unescape(str(text))
    text = re.sub(r"[\r\n]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip().lower()


_ALIAS_TO_STANDARD: Dict[str, str] = {}
for _header in STANDARD_HEADERS:
    _ALIAS_TO_STANDARD[_normalize_header(_header)] = _header
    for _alias in _EXTRA_ALIASES.get(_header, []):
        _ALIAS_TO_STANDARD[_normalize_header(_alias)] = _header


def _headers_in_row(ws, row_idx: int) -> Dict[str, int]:
    """Map each recognised standard header in this row to its column number
    (the first one wins if a header appears twice)."""
    found: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        standard = _ALIAS_TO_STANDARD.get(_normalize_header(ws.cell(row=row_idx, column=col_idx).value))
        if standard and standard not in found:
            found[standard] = col_idx
    return found


def _is_header_row(found: Dict[str, int]) -> bool:
    return "Risk ID" in found and ("Control ID" in found or "Control Reference" in found)


def _find_header_row(ws) -> Optional[int]:
    """The header row is the first of the top 10 rows that names both a risk
    and a control identifier column."""
    for row_idx in range(1, min(11, ws.max_row + 1)):
        if _is_header_row(_headers_in_row(ws, row_idx)):
            return row_idx
    return None


def _cell(ws, row: int, col_idx: Optional[int], keep_newlines: bool = False) -> Optional[str]:
    if col_idx is None:
        return None
    value = ws.cell(row=row, column=col_idx).value
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in _NO_VALUE:
        return None
    if not keep_newlines:
        text = re.sub(r"\s+", " ", text)
    return text or None


def _pick_worksheet(wb):
    visible = [wb[name] for name in wb.sheetnames if wb[name].sheet_state != "hidden"]
    for sheet in visible:
        if sheet.title.lower() == "racm upload":
            return sheet
    return visible[0] if visible else None


def extract_racm_from_excel(file_bytes: bytes) -> Dict[str, Any]:
    """Read the first control row of a RACM spreadsheet.

    Returns a dict shaped like RACMControlRecord. Raises ExcelParsingError for
    an unreadable file or missing headers, and ExcelValidationError (with
    per-row details) when a data row lacks a required identifier.
    """
    if not openpyxl:
        raise ExcelParsingError("openpyxl not installed")

    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ExcelParsingError(f"Failed to load workbook: {str(e)}")

    ws = _pick_worksheet(wb)
    if ws is None:
        raise ExcelParsingError("No usable worksheet found")
    if ws.max_row < 2:
        raise ExcelParsingError("Worksheet is empty")

    header_row = _find_header_row(ws)
    if header_row is None:
        raise ExcelParsingError(
            "Cannot detect header row - no row names both a risk column (Risk ID) "
            "and a control column (Control ID or Control Reference)"
        )
    cols = _headers_in_row(ws, header_row)

    errors: List[Dict[str, Any]] = []
    extracted: List[Dict[str, Any]] = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        if all(ws.cell(row=row_idx, column=c).value is None for c in range(1, ws.max_column + 1)):
            continue

        def get(header: str) -> Optional[str]:
            return _cell(ws, row_idx, cols.get(header), keep_newlines=header in _MULTILINE)

        # A file with only one control column uses it as both the id and the code.
        control_reference = get("Control Reference") or get("Control ID")
        control_id = get("Control ID") or get("Control Reference")
        risk_id = get("Risk ID")

        row_errors = []
        if not control_id:
            row_errors.append({"field": "control.control_id", "message": "Control ID (or Control Reference) is required"})
        if not risk_id:
            row_errors.append({"field": "control.risk.risk_id", "message": "Risk ID is required"})
        if row_errors:
            errors.extend({"sheet": ws.title, "row": row_idx, **e} for e in row_errors)
            continue

        extracted.append(
            {
                "schema_version": "1.0",
                "engagement": {
                    "engagement_id": get("Engagement ID") or f"{ws.title}_{row_idx}".replace(" ", "_"),
                    "client": get("Client") or "Not specified",
                    "opco": get("OpCo"),
                    "area": get("Area"),
                    "mega_process": get("Mega Process"),
                    "process": get("Process") or "Unknown Process",
                    "subprocess": get("Sub-process"),
                },
                "control": {
                    "control_id": control_id,
                    "control_title": get("Control Title") or control_id,
                    "control_reference": control_reference,
                    "risk": {
                        "risk_id": risk_id,
                        "risk_description": get("Risk Description") or risk_id,
                    },
                    "control_objective": get("Control Objective"),
                    "suggested_control_description": get("Suggested Control Description"),
                    "control_frequency": get("Control Frequency"),
                    "control_nature": get("Control Nature"),
                    "control_type": get("Control Type"),
                    "expected_evidence": get("Expected Evidence"),
                    "suggested_test_procedures_tod": get("Suggested Test Procedures (ToD)"),
                    "suggested_test_procedures_toe": get("Suggested Test Procedures (ToE)"),
                },
            }
        )

    if errors:
        raise ExcelValidationError(errors)
    if not extracted:
        raise ExcelParsingError("No valid data rows found in worksheet")

    # The app works on one control at a time, so only the first row is used.
    return extracted[0]
