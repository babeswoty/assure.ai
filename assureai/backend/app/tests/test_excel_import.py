"""Tests for Excel RACM import functionality."""

import pytest
from io import BytesIO

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.excel_import import extract_racm_from_excel, ExcelParsingError, ExcelValidationError


def create_test_workbook(
    headers: dict[str, int],
    rows: list[dict[str, str]],
    sheet_name: str = "RACM Upload"
) -> bytes:
    """Helper to create test Excel workbooks.

    Args:
        headers: Mapping of header names to column indices (1-based)
        rows: List of row data dicts
        sheet_name: Name of the worksheet

    Returns:
        Excel file bytes
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Write headers
    for header_name, col_idx in headers.items():
        ws.cell(row=1, column=col_idx, value=header_name)

    # Write data rows
    for row_idx, row_data in enumerate(rows, start=2):
        for header_name, col_idx in headers.items():
            value = row_data.get(header_name)
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Write to bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


class TestExcelImport:
    """Basic Excel import tests."""

    def test_minimal_valid_racm(self):
        """Test importing minimal valid RACM with required fields."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Process Name": 3,
            "Risk Description": 4,
        }
        rows = [
            {
                "Control Ref": "A.C.C.1",
                "Risk Ref": "A.C.R.1",
                "Process Name": "User Access Management",
                "Risk Description": "Unauthorized access",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)
        result = extract_racm_from_excel(file_bytes)

        assert result["control"]["control_id"] == "A.C.C.1"
        assert result["control"]["risk"]["risk_id"] == "A.C.R.1"
        assert result["engagement"]["process"] == "User Access Management"
        assert result["control"]["risk"]["risk_description"] == "Unauthorized access"

    def test_all_fields_populated(self):
        """Test importing all available RACM fields."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Mega Process": 3,
            "Major Process Name": 4,
            "Process Name": 5,
            "Sub-process": 6,
            "Control Objective": 7,
            "Risk Description": 8,
            "AML/CTF Control": 9,
            "Suggested Control": 10,
            "Control Nature": 11,
            "Control Type": 12,
            "Frequency of Control": 13,
            "Control Evidence": 14,
        }
        rows = [
            {
                "Control Ref": "CTRL-001",
                "Risk Ref": "RISK-001",
                "Mega Process": "IT Controls",
                "Major Process Name": "Access Management",
                "Process Name": "User Provisioning",
                "Sub-process": "Onboarding",
                "Control Objective": "Ensure authorized access",
                "Risk Description": "Unauthorized system access",
                "AML/CTF Control": "Access Control",
                "Suggested Control": "Review and approve all access requests",
                "Control Nature": "Detective",
                "Control Type": "Automated",
                "Frequency of Control": "Monthly",
                "Control Evidence": "Access logs and approval records",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)
        result = extract_racm_from_excel(file_bytes)

        assert result["control"]["control_id"] == "CTRL-001"
        assert result["control"]["control_title"] == "Access Control"
        assert result["engagement"]["mega_process"] == "IT Controls"
        assert result["engagement"]["area"] == "Access Management"
        assert result["control"]["suggested_control_description"] == "Review and approve all access requests"

    def test_multiline_cells_preserved(self):
        """Test that multiline cells preserve line breaks."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Process Name": 3,
            "Risk Description": 4,
            "Suggested Control": 5,
            "Control Evidence": 6,
        }
        rows = [
            {
                "Control Ref": "A.C.C.1",
                "Risk Ref": "A.C.R.1",
                "Process Name": "Process",
                "Risk Description": "Risk",
                "Suggested Control": "Step 1\nStep 2\nStep 3",
                "Control Evidence": "Point 1\nPoint 2\nPoint 3",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)
        result = extract_racm_from_excel(file_bytes)

        # Line breaks should be preserved in multiline fields
        assert "\n" in result["control"]["suggested_control_description"]
        assert "\n" in result["control"]["expected_evidence"]

    def test_header_aliases(self):
        """Test that header aliases work."""
        # Use non-standard header names (aliases)
        headers = {
            "control ref (#)": 1,  # Alias for "Control Ref"
            "risk ref (#)": 2,  # Alias for "Risk Ref"
            "process name": 3,
            "risk description": 4,
        }
        rows = [
            {
                "control ref (#)": "A.C.C.1",
                "risk ref (#)": "A.C.R.1",
                "process name": "Process",
                "risk description": "Risk",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)
        result = extract_racm_from_excel(file_bytes)

        assert result["control"]["control_id"] == "A.C.C.1"
        assert result["control"]["risk"]["risk_id"] == "A.C.R.1"

    def test_blank_optional_fields(self):
        """Test that blank optional fields are handled correctly."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Process Name": 3,
            "Risk Description": 4,
            "Control Objective": 5,
            "Control Nature": 6,
        }
        rows = [
            {
                "Control Ref": "A.C.C.1",
                "Risk Ref": "A.C.R.1",
                "Process Name": "Process",
                "Risk Description": "Risk",
                # Blank optional fields
            }
        ]

        file_bytes = create_test_workbook(headers, rows)
        result = extract_racm_from_excel(file_bytes)

        assert result["control"]["control_objective"] is None
        assert result["control"]["control_nature"] is None

    def test_missing_control_ref_fails(self):
        """Test that missing Control Ref returns 422 validation error."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Process Name": 3,
            "Risk Description": 4,
        }
        rows = [
            {
                # Missing Control Ref
                "Risk Ref": "A.C.R.1",
                "Process Name": "Process",
                "Risk Description": "Risk",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)

        with pytest.raises(ExcelValidationError) as exc_info:
            extract_racm_from_excel(file_bytes)

        assert len(exc_info.value.errors) > 0
        assert any(e["field"] == "control.control_id" for e in exc_info.value.errors)

    def test_missing_risk_ref_fails(self):
        """Test that missing Risk Ref returns 422 validation error."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Process Name": 3,
            "Risk Description": 4,
        }
        rows = [
            {
                "Control Ref": "A.C.C.1",
                # Missing Risk Ref
                "Process Name": "Process",
                "Risk Description": "Risk",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)

        with pytest.raises(ExcelValidationError) as exc_info:
            extract_racm_from_excel(file_bytes)

        assert len(exc_info.value.errors) > 0
        assert any(e["field"] == "control.risk.risk_id" for e in exc_info.value.errors)

    def test_empty_workbook_fails(self):
        """Test that empty workbook returns 400 error."""
        wb = openpyxl.Workbook()
        ws = wb.active

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        file_bytes = output.getvalue()

        with pytest.raises(ExcelParsingError):
            extract_racm_from_excel(file_bytes)

    def test_unsupported_file_type(self):
        """Test that unsupported file types are rejected at API level (not here)."""
        # This is handled by the API endpoint, not the parser
        pass

    def test_default_values_applied(self):
        """Test that safe defaults are applied for missing engagement fields."""
        headers = {
            "Control Ref": 1,
            "Risk Ref": 2,
            "Process Name": 3,
            "Risk Description": 4,
        }
        rows = [
            {
                "Control Ref": "A.C.C.1",
                "Risk Ref": "A.C.R.1",
                "Process Name": "Process",
                "Risk Description": "Risk",
            }
        ]

        file_bytes = create_test_workbook(headers, rows)
        result = extract_racm_from_excel(file_bytes)

        # Check defaults are applied
        assert result["engagement"]["client"] == "MTN Fintech"
        assert result["engagement"]["opco"] == "Group Fintech / Fintech OpCo"
        assert result["engagement"]["process"] == "Process"

    def test_racm_upload_worksheet_preferred(self):
        """Test that RACM Upload worksheet is used if it exists."""
        # Create workbook with multiple sheets
        wb = openpyxl.Workbook()
        wb.create_sheet("Read Me")
        racm_sheet = wb.create_sheet("RACM Upload")

        # Add headers and data to RACM Upload sheet
        racm_sheet.cell(row=1, column=1, value="Control Ref")
        racm_sheet.cell(row=1, column=2, value="Risk Ref")
        racm_sheet.cell(row=1, column=3, value="Process Name")
        racm_sheet.cell(row=1, column=4, value="Risk Description")

        racm_sheet.cell(row=2, column=1, value="A.C.C.1")
        racm_sheet.cell(row=2, column=2, value="A.C.R.1")
        racm_sheet.cell(row=2, column=3, value="Process")
        racm_sheet.cell(row=2, column=4, value="Risk")

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        file_bytes = output.getvalue()

        result = extract_racm_from_excel(file_bytes)
        assert result["control"]["control_id"] == "A.C.C.1"

    def test_header_detection_across_rows(self):
        """Test that header row is detected even if not in row 1."""
        wb = openpyxl.Workbook()
        ws = wb.active

        # Put some junk in rows 1-2
        ws.cell(row=1, column=1, value="Title Row")
        ws.cell(row=2, column=1, value="Subtitle Row")

        # Put actual headers in row 3
        ws.cell(row=3, column=1, value="Control Ref")
        ws.cell(row=3, column=2, value="Risk Ref")
        ws.cell(row=3, column=3, value="Process Name")
        ws.cell(row=3, column=4, value="Risk Description")

        # Put data in row 4
        ws.cell(row=4, column=1, value="A.C.C.1")
        ws.cell(row=4, column=2, value="A.C.R.1")
        ws.cell(row=4, column=3, value="Process")
        ws.cell(row=4, column=4, value="Risk")

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        file_bytes = output.getvalue()

        result = extract_racm_from_excel(file_bytes)
        assert result["control"]["control_id"] == "A.C.C.1"
